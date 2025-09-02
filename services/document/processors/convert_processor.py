'''
Description: 文档格式转换处理器
Author: zyq
Date: 2025-09-01 16:51:00
LastEditors: zyq
LastEditTime: 2025-09-02 10:46:32
'''

import os
import tempfile
import asyncio
from typing import Dict, Any
from datetime import datetime
from concurrent.futures import ProcessPoolExecutor
from loguru import logger

from core.schemas.file_models import FileInfo
from ..schemas import DocumentConvertRequest, DocumentTaskTypePrefix
from ..utils.file_manager import FileManager
from .base_processor import BaseProcessor


# WeasyPrint异步执行函数（模块级别避免pickle错误）
def _weasyprint_write_pdf(html_content: str, output_path: str):
    """在独立进程中执行WeasyPrint PDF生成"""
    import weasyprint
    weasyprint.HTML(string=html_content).write_pdf(output_path)


# 全局共享的进程池和并发控制
_pdf_executor = ProcessPoolExecutor(max_workers=2)
_pdf_semaphore = asyncio.Semaphore(2)


class ConvertProcessor(BaseProcessor):
    """文档格式转换处理器"""
    
    def __init__(self, config: Dict[str, Any]):
        """初始化转换处理器"""
        super().__init__(config)
        self.file_manager = FileManager()
        
        # 支持的转换类型
        self.supported_conversions = {
            ('docx', 'pdf'): self._convert_docx_to_pdf,
            ('docx', 'markdown'): self._convert_docx_to_markdown,
            ('xlsx', 'pdf'): self._convert_xlsx_to_pdf,
            ('xlsx', 'docx'): self._convert_xlsx_to_docx,
            ('xlsx', 'markdown'): self._convert_xlsx_to_markdown,
            ('markdown', 'pdf'): self._convert_markdown_to_pdf,
            ('markdown', 'docx'): self._convert_markdown_to_docx,
        }
    
    async def process(self, request: DocumentConvertRequest, temp_file_path: str, file_info: FileInfo, trace_id: str = None, task_id: str = None) -> str:
        """
        执行文档格式转换
        
        Args:
            request: 转换请求
            temp_file_path: 临时文件路径
            file_info: 文件信息
            trace_id: 链路追踪ID
            
        Returns:
            转换后文件的存储路径
        """
        self._log_processing_start("ConvertProcessor", f"{request.source_format}→{request.target_format}", trace_id)
        
        try:
            # 检查转换类型是否支持
            convert_key = (request.source_format, request.target_format)
            convert_func = self.supported_conversions.get(convert_key)
            
            if not convert_func:
                raise ValueError(f"不支持的转换类型: {request.source_format} → {request.target_format}")
            
            # 执行转换
            converted_file_path = await convert_func(temp_file_path, file_info, request, trace_id, task_id)
            
            self._log_processing_end("ConvertProcessor", f"{request.source_format}→{request.target_format}", trace_id)
            return converted_file_path
            
        except Exception as e:
            self._log_error("ConvertProcessor", trace_id, e)
            self._log_processing_end("ConvertProcessor", f"{request.source_format}→{request.target_format}", trace_id, success=False)
            raise e
    
    # ==================== DOCX转换实现 ====================
    
    async def _convert_docx_to_pdf(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str) -> str:
        """DOCX转PDF转换"""
        try:
            # 高级转换：python-docx → HTML → weasyprint
            return await self._docx_to_pdf_advanced(source_path, file_info, request, trace_id, task_id)
        except Exception as e:
            logger.warning(f"DOCX→PDF高级转换失败: {e}")
            try:
                # 简化转换：基础HTML → weasyprint
                return await self._docx_to_pdf_simple(source_path, file_info, request, trace_id, task_id)
            except Exception as e2:
                logger.warning(f"DOCX→PDF简化转换失败: {e2}")
                # 基础转换：纯文本PDF
                return await self._docx_to_pdf_fallback(source_path, file_info, request, trace_id, task_id)
    
    async def _docx_to_pdf_advanced(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str) -> str:
        """DOCX→PDF高级转换：保持完整格式"""
        try:
            from docx import Document
            import weasyprint
            
            # 读取DOCX文档
            doc = Document(source_path)
            
            # 转换为HTML
            html_content = self._docx_to_html_advanced(doc)
            
            # 生成输出文件路径
            output_filename = f"{request.filename.rsplit('.', 1)[0]}.pdf"
            output_path = self.file_manager.create_file_path(task_id, output_filename)
            
            # 使用weasyprint生成PDF
            # 使用全局进程池和信号量控制并发
            async with _pdf_semaphore:
                loop = asyncio.get_event_loop()
                await loop.run_in_executor(
                    _pdf_executor, 
                    _weasyprint_write_pdf, 
                    html_content, 
                    output_path
                )
            
            # 注册文件到文件管理器
            self.file_manager.register_file(task_id, output_filename, output_path, expire_hours=24)
            
            logger.info(f"DOCX→PDF高级转换成功: {output_path}")
            return output_path
            
        except ImportError as e:
            raise Exception(f"缺少依赖库: {str(e)}")
        except Exception as e:
            raise Exception(f"高级转换失败: {str(e)}")
    
    async def _docx_to_pdf_simple(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str) -> str:
        """DOCX→PDF简化转换：基础格式"""
        try:
            from docx import Document
            import weasyprint
            
            # 读取DOCX文档
            doc = Document(source_path)
            
            # 转换为简单HTML
            html_content = self._docx_to_html_simple(doc)
            
            # 生成输出文件路径
            output_filename = f"{request.filename.rsplit('.', 1)[0]}.pdf"
            output_path = self.file_manager.create_file_path(task_id, output_filename)
            
            # 使用weasyprint生成PDF
            # 使用全局进程池和信号量控制并发
            async with _pdf_semaphore:
                loop = asyncio.get_event_loop()
                await loop.run_in_executor(
                    _pdf_executor, 
                    _weasyprint_write_pdf, 
                    html_content, 
                    output_path
                )
            
            # 注册文件到文件管理器
            self.file_manager.register_file(task_id, output_filename, output_path, expire_hours=24)
            
            logger.info(f"DOCX→PDF简化转换成功: {output_path}")
            return output_path
            
        except Exception as e:
            raise Exception(f"简化转换失败: {str(e)}")
    
    async def _docx_to_pdf_fallback(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str) -> str:
        """DOCX→PDF基础转换：纯文本PDF"""
        try:
            from docx import Document
            import weasyprint
            
            # 读取DOCX文档，只提取纯文本
            doc = Document(source_path)
            text_content = "\n".join([paragraph.text for paragraph in doc.paragraphs])
            
            # 生成简单HTML
            html_content = f"""
            <!DOCTYPE html>
            <html>
            <head>
                <meta charset="utf-8">
                <title>{request.filename}</title>
                <style>
                    body {{ font-family: Arial, sans-serif; margin: 40px; line-height: 1.6; }}
                    h1 {{ color: #333; }}
                    pre {{ white-space: pre-wrap; }}
                </style>
            </head>
            <body>
                <h1>Document Conversion Result</h1>
                <p><strong>Source:</strong> {request.filename}</p>
                <p><strong>Converted:</strong> {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p>
                <hr>
                <pre>{text_content}</pre>
            </body>
            </html>
            """
            
            # 生成输出文件路径
            output_filename = f"{request.filename.rsplit('.', 1)[0]}.pdf"
            output_path = self.file_manager.create_file_path(task_id, output_filename)
            
            # 使用weasyprint生成PDF
            # 使用全局进程池和信号量控制并发
            async with _pdf_semaphore:
                loop = asyncio.get_event_loop()
                await loop.run_in_executor(
                    _pdf_executor, 
                    _weasyprint_write_pdf, 
                    html_content, 
                    output_path
                )
            
            # 注册文件到文件管理器
            self.file_manager.register_file(task_id, output_filename, output_path, expire_hours=24)
            
            logger.info(f"DOCX→PDF基础转换成功: {output_path}")
            return output_path
            
        except Exception as e:
            logger.error(f"基础转换失败: {str(e)}")
            raise Exception("所有转换方法失败")
    
    async def _convert_docx_to_markdown(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str) -> str:
        """DOCX转Markdown转换"""
        try:
            # 高级转换：保持格式
            return await self._docx_to_markdown_advanced(source_path, file_info, request, trace_id, task_id)
        except Exception as e:
            logger.warning(f"DOCX→Markdown高级转换失败: {e}")
            try:
                # 简化转换：基础格式
                return await self._docx_to_markdown_simple(source_path, file_info, request, trace_id, task_id)
            except Exception as e2:
                logger.warning(f"DOCX→Markdown简化转换失败: {e2}")
                # 基础转换：纯文本
                return await self._docx_to_markdown_fallback(source_path, file_info, request, trace_id, task_id)
    
    async def _docx_to_markdown_advanced(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str) -> str:
        """DOCX转Markdown高级实现"""
        from docx import Document
        
        doc = Document(source_path)
        markdown_parts = []
        
        # 转换段落和图片
        for paragraph in doc.paragraphs:
            # 提取段落中的图片
            images_md = self._extract_paragraph_images_to_markdown(paragraph)
            
            if paragraph.text.strip():
                if paragraph.style.name.startswith('Heading'):
                    level = min(int(paragraph.style.name[-1]) if paragraph.style.name[-1].isdigit() else 1, 6)
                    markdown_parts.append(f"{'#' * level} {paragraph.text}")
                else:
                    # 处理段落格式
                    md_text = self._process_paragraph_to_markdown(paragraph)
                    markdown_parts.append(md_text)
            
            # 添加图片
            if images_md:
                markdown_parts.extend(images_md)
        
        # 转换表格
        for table in doc.tables:
            markdown_parts.append(self._table_to_markdown(table))
        
        # 保存Markdown文件
        output_filename = f"{request.filename.rsplit('.', 1)[0]}.md"
        output_path = self.file_manager.create_file_path(task_id, output_filename)
        
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write('\n\n'.join(markdown_parts))
        
        # 注册文件到文件管理器
        self.file_manager.register_file(task_id, output_filename, output_path, expire_hours=24)
        
        return output_path
    
    async def _docx_to_markdown_simple(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str) -> str:
        """DOCX转Markdown简化实现"""
        from docx import Document
        
        doc = Document(source_path)
        markdown_parts = []
        
        # 只转换基础文本
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                markdown_parts.append(paragraph.text)
        
        # 保存Markdown文件
        output_filename = f"{request.filename.rsplit('.', 1)[0]}.md"
        output_path = self.file_manager.create_file_path(task_id, output_filename)
        
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write('\n\n'.join(markdown_parts))
        
        # 注册文件到文件管理器
        self.file_manager.register_file(task_id, output_filename, output_path, expire_hours=24)
        
        return output_path
    
    async def _docx_to_markdown_fallback(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str) -> str:
        """DOCX转Markdown基础实现"""
        return await self._docx_to_markdown_simple(source_path, file_info, request, trace_id, task_id)
    
    # ==================== XLSX转换实现 ====================
    
    async def _convert_xlsx_to_pdf(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str) -> str:
        """XLSX转PDF转换"""
        try:
            return await self._xlsx_to_pdf_advanced(source_path, file_info, request, trace_id, task_id)
        except Exception as e:
            logger.warning(f"XLSX→PDF高级转换失败: {e}")
            try:
                return await self._xlsx_to_pdf_simple(source_path, file_info, request, trace_id, task_id)
            except Exception as e2:
                logger.warning(f"XLSX→PDF简化转换失败: {e2}")
                return await self._xlsx_to_pdf_fallback(source_path, file_info, request, trace_id, task_id)
    
    async def _convert_xlsx_to_docx(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str = None) -> str:
        """XLSX转DOCX转换"""
        try:
            return await self._xlsx_to_docx_advanced(source_path, file_info, request, trace_id, task_id)
        except Exception as e:
            logger.warning(f"XLSX→DOCX高级转换失败: {e}")
            try:
                return await self._xlsx_to_docx_simple(source_path, file_info, request, trace_id, task_id)
            except Exception as e2:
                logger.warning(f"XLSX→DOCX简化转换失败: {e2}")
                return await self._xlsx_to_docx_fallback(source_path, file_info, request, trace_id, task_id)
    
    async def _convert_xlsx_to_markdown(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str = None) -> str:
        """XLSX转Markdown转换"""
        try:
            return await self._xlsx_to_markdown_advanced(source_path, file_info, request, trace_id, task_id)
        except Exception as e:
            logger.warning(f"XLSX→Markdown高级转换失败: {e}")
            try:
                return await self._xlsx_to_markdown_simple(source_path, file_info, request, trace_id, task_id)
            except Exception as e2:
                logger.warning(f"XLSX→Markdown简化转换失败: {e2}")
                return await self._xlsx_to_markdown_fallback(source_path, file_info, request, trace_id, task_id)
    
    async def _xlsx_to_pdf_advanced(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str = None) -> str:
        """XLSX转PDF高级实现：保持完整格式"""
        try:
            from openpyxl import load_workbook
            import weasyprint
            
            # 读取Excel文件
            wb = load_workbook(source_path)
            
            # 转换为HTML
            html_content = self._xlsx_to_html_advanced(wb)
            
            # 生成输出文件路径
            task_id = task_id or self.file_manager.generate_task_id(DocumentTaskTypePrefix.CONVERT_TASK.value)
            output_filename = f"{request.filename.rsplit('.', 1)[0]}.pdf"
            output_path = self.file_manager.create_file_path(task_id, output_filename)
            
            # 使用weasyprint生成PDF
            # 使用全局进程池和信号量控制并发
            async with _pdf_semaphore:
                loop = asyncio.get_event_loop()
                await loop.run_in_executor(
                    _pdf_executor, 
                    _weasyprint_write_pdf, 
                    html_content, 
                    output_path
                )
            
            # 注册文件到文件管理器
            self.file_manager.register_file(task_id, output_filename, output_path, expire_hours=24)
            
            logger.info(f"XLSX→PDF高级转换成功: {output_path}")
            return output_path
            
        except Exception as e:
            raise Exception(f"高级转换失败: {str(e)}")
    
    async def _xlsx_to_pdf_simple(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str = None) -> str:
        """XLSX转PDF简化实现：基础表格格式"""
        try:
            from openpyxl import load_workbook
            import weasyprint
            
            # 读取Excel文件
            wb = load_workbook(source_path)
            
            # 转换为简单HTML
            html_content = self._xlsx_to_html_simple(wb)
            
            # 生成输出文件路径
            task_id = task_id or self.file_manager.generate_task_id(DocumentTaskTypePrefix.CONVERT_TASK.value)
            output_filename = f"{request.filename.rsplit('.', 1)[0]}.pdf"
            output_path = self.file_manager.create_file_path(task_id, output_filename)
            
            # 使用weasyprint生成PDF
            # 使用全局进程池和信号量控制并发
            async with _pdf_semaphore:
                loop = asyncio.get_event_loop()
                await loop.run_in_executor(
                    _pdf_executor, 
                    _weasyprint_write_pdf, 
                    html_content, 
                    output_path
                )
            
            # 注册文件到文件管理器
            self.file_manager.register_file(task_id, output_filename, output_path, expire_hours=24)
            
            logger.info(f"XLSX→PDF简化转换成功: {output_path}")
            return output_path
            
        except Exception as e:
            raise Exception(f"简化转换失败: {str(e)}")
    
    async def _xlsx_to_pdf_fallback(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str = None) -> str:
        """XLSX转PDF基础实现：纯文本PDF"""
        try:
            from openpyxl import load_workbook
            import weasyprint
            
            # 读取Excel文件，只提取文本内容
            wb = load_workbook(source_path)
            text_content = self._xlsx_to_text(wb)
            
            # 生成简单HTML
            html_content = f"""
            <!DOCTYPE html>
            <html>
            <head>
                <meta charset="utf-8">
                <title>{request.filename}</title>
                <style>
                    body {{ font-family: Arial, sans-serif; margin: 40px; line-height: 1.6; }}
                    h1 {{ color: #333; }}
                    pre {{ white-space: pre-wrap; }}
                </style>
            </head>
            <body>
                <h1>Excel Data Export</h1>
                <p><strong>Source:</strong> {request.filename}</p>
                <p><strong>Converted:</strong> {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p>
                <hr>
                <pre>{text_content}</pre>
            </body>
            </html>
            """
            
            # 生成输出文件路径
            task_id = task_id or self.file_manager.generate_task_id(DocumentTaskTypePrefix.CONVERT_TASK.value)
            output_filename = f"{request.filename.rsplit('.', 1)[0]}.pdf"
            output_path = self.file_manager.create_file_path(task_id, output_filename)
            
            # 使用weasyprint生成PDF
            # 使用全局进程池和信号量控制并发
            async with _pdf_semaphore:
                loop = asyncio.get_event_loop()
                await loop.run_in_executor(
                    _pdf_executor, 
                    _weasyprint_write_pdf, 
                    html_content, 
                    output_path
                )
            
            # 注册文件到文件管理器
            self.file_manager.register_file(task_id, output_filename, output_path, expire_hours=24)
            
            logger.info(f"XLSX→PDF基础转换成功: {output_path}")
            return output_path
            
        except Exception as e:
            logger.error(f"基础转换失败: {str(e)}")
            raise Exception("所有转换方法失败")
    
    async def _xlsx_to_docx_advanced(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str = None) -> str:
        """XLSX转DOCX高级实现：保持表格格式"""
        try:
            from openpyxl import load_workbook
            from docx import Document
            from docx.shared import Inches
            
            # 读取Excel文件
            wb = load_workbook(source_path)
            
            # 创建Word文档
            doc = Document()
            doc.add_heading(f'Excel Data: {request.filename}', 0)
            
            # 转换每个工作表
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                
                # 添加工作表标题
                doc.add_heading(f'Sheet: {sheet_name}', level=1)
                
                # 获取数据范围
                data_rows = list(ws.iter_rows(values_only=True))
                if not data_rows:
                    doc.add_paragraph('No data in this sheet')
                    continue
                
                # 创建表格
                table = doc.add_table(rows=len(data_rows), cols=len(data_rows[0]))
                table.style = 'Table Grid'
                
                # 填充表格数据
                for i, row_data in enumerate(data_rows):
                    for j, cell_value in enumerate(row_data):
                        if cell_value is not None:
                            table.cell(i, j).text = str(cell_value)
                
                doc.add_paragraph()  # 添加段落间隔
            
            # 生成输出文件路径
            task_id = task_id or self.file_manager.generate_task_id(DocumentTaskTypePrefix.CONVERT_TASK.value)
            output_filename = f"{request.filename.rsplit('.', 1)[0]}.docx"
            output_path = self.file_manager.create_file_path(task_id, output_filename)
            
            # 保存文档
            doc.save(output_path)
            
            # 注册文件到文件管理器
            self.file_manager.register_file(task_id, output_filename, output_path, expire_hours=24)
            
            logger.info(f"XLSX→DOCX高级转换成功: {output_path}")
            return output_path
            
        except Exception as e:
            raise Exception(f"高级转换失败: {str(e)}")
    
    async def _xlsx_to_docx_simple(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str = None) -> str:
        """XLSX转DOCX简化实现：基础表格"""
        try:
            from openpyxl import load_workbook
            from docx import Document
            
            # 读取Excel文件
            wb = load_workbook(source_path)
            
            # 创建Word文档
            doc = Document()
            doc.add_heading(f'Excel Data: {request.filename}', 0)
            
            # 只处理第一个工作表
            ws = wb.active
            data_rows = list(ws.iter_rows(values_only=True))
            
            if data_rows:
                # 创建表格
                table = doc.add_table(rows=len(data_rows), cols=len(data_rows[0]))
                
                # 填充表格数据
                for i, row_data in enumerate(data_rows):
                    for j, cell_value in enumerate(row_data):
                        if cell_value is not None:
                            table.cell(i, j).text = str(cell_value)
            
            # 生成输出文件路径
            task_id = task_id or self.file_manager.generate_task_id(DocumentTaskTypePrefix.CONVERT_TASK.value)
            output_filename = f"{request.filename.rsplit('.', 1)[0]}.docx"
            output_path = self.file_manager.create_file_path(task_id, output_filename)
            
            # 保存文档
            doc.save(output_path)
            
            # 注册文件到文件管理器
            self.file_manager.register_file(task_id, output_filename, output_path, expire_hours=24)
            
            logger.info(f"XLSX→DOCX简化转换成功: {output_path}")
            return output_path
            
        except Exception as e:
            raise Exception(f"简化转换失败: {str(e)}")
    
    async def _xlsx_to_docx_fallback(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str = None) -> str:
        """XLSX转DOCX基础实现：纯文本"""
        try:
            from openpyxl import load_workbook
            from docx import Document
            
            # 读取Excel文件，只提取文本内容
            wb = load_workbook(source_path)
            text_content = self._xlsx_to_text(wb)
            
            # 创建Word文档
            doc = Document()
            doc.add_heading(f'Excel Data: {request.filename}', 0)
            doc.add_paragraph(f'Converted: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')
            doc.add_paragraph(text_content)
            
            # 生成输出文件路径
            task_id = task_id or self.file_manager.generate_task_id(DocumentTaskTypePrefix.CONVERT_TASK.value)
            output_filename = f"{request.filename.rsplit('.', 1)[0]}.docx"
            output_path = self.file_manager.create_file_path(task_id, output_filename)
            
            # 保存文档
            doc.save(output_path)
            
            # 注册文件到文件管理器
            self.file_manager.register_file(task_id, output_filename, output_path, expire_hours=24)
            
            logger.info(f"XLSX→DOCX基础转换成功: {output_path}")
            return output_path
            
        except Exception as e:
            logger.error(f"基础转换失败: {str(e)}")
            raise Exception("所有转换方法失败")
    
    async def _xlsx_to_markdown_advanced(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str = None) -> str:
        """XLSX转Markdown高级实现：完整表格格式"""
        try:
            from openpyxl import load_workbook
            
            # 读取Excel文件
            wb = load_workbook(source_path)
            markdown_parts = []
            
            markdown_parts.append(f"# Excel Data: {request.filename}")
            markdown_parts.append(f"**Converted:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
            markdown_parts.append("")
            
            # 转换每个工作表
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                
                markdown_parts.append(f"## Sheet: {sheet_name}")
                markdown_parts.append("")
                
                # 获取数据范围
                data_rows = list(ws.iter_rows(values_only=True))
                if not data_rows:
                    markdown_parts.append("No data in this sheet")
                    markdown_parts.append("")
                    continue
                
                # 转换为Markdown表格
                markdown_table = self._excel_data_to_markdown_table(data_rows)
                markdown_parts.append(markdown_table)
                markdown_parts.append("")
            
            # 保存Markdown文件
            task_id = task_id or self.file_manager.generate_task_id(DocumentTaskTypePrefix.CONVERT_TASK.value)
            output_filename = f"{request.filename.rsplit('.', 1)[0]}.md"
            output_path = self.file_manager.create_file_path(task_id, output_filename)
            
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write('\n'.join(markdown_parts))
            
            # 注册文件到文件管理器
            self.file_manager.register_file(task_id, output_filename, output_path, expire_hours=24)
            
            logger.info(f"XLSX→Markdown高级转换成功: {output_path}")
            return output_path
            
        except Exception as e:
            raise Exception(f"高级转换失败: {str(e)}")
    
    async def _xlsx_to_markdown_simple(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str = None) -> str:
        """XLSX转Markdown简化实现：基础表格"""
        try:
            from openpyxl import load_workbook
            
            # 读取Excel文件
            wb = load_workbook(source_path)
            
            # 只处理第一个工作表
            ws = wb.active
            data_rows = list(ws.iter_rows(values_only=True))
            
            markdown_parts = []
            markdown_parts.append(f"# {request.filename}")
            markdown_parts.append("")
            
            if data_rows:
                # 转换为Markdown表格
                markdown_table = self._excel_data_to_markdown_table(data_rows)
                markdown_parts.append(markdown_table)
            
            # 保存Markdown文件
            task_id = task_id or self.file_manager.generate_task_id(DocumentTaskTypePrefix.CONVERT_TASK.value)
            output_filename = f"{request.filename.rsplit('.', 1)[0]}.md"
            output_path = self.file_manager.create_file_path(task_id, output_filename)
            
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write('\n'.join(markdown_parts))
            
            # 注册文件到文件管理器
            self.file_manager.register_file(task_id, output_filename, output_path, expire_hours=24)
            
            logger.info(f"XLSX→Markdown简化转换成功: {output_path}")
            return output_path
            
        except Exception as e:
            raise Exception(f"简化转换失败: {str(e)}")
    
    async def _xlsx_to_markdown_fallback(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str = None) -> str:
        """XLSX转Markdown基础实现：纯文本"""
        try:
            from openpyxl import load_workbook
            
            # 读取Excel文件，只提取文本内容
            wb = load_workbook(source_path)
            text_content = self._xlsx_to_text(wb)
            
            # 生成简单Markdown
            markdown_content = f"""# {request.filename}

**Converted:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}

## Data

```
{text_content}
```
"""
            
            # 保存Markdown文件
            task_id = task_id or self.file_manager.generate_task_id(DocumentTaskTypePrefix.CONVERT_TASK.value)
            output_filename = f"{request.filename.rsplit('.', 1)[0]}.md"
            output_path = self.file_manager.create_file_path(task_id, output_filename)
            
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(markdown_content)
            
            # 注册文件到文件管理器
            self.file_manager.register_file(task_id, output_filename, output_path, expire_hours=24)
            
            logger.info(f"XLSX→Markdown基础转换成功: {output_path}")
            return output_path
            
        except Exception as e:
            logger.error(f"基础转换失败: {str(e)}")
            raise Exception("所有转换方法失败")
    
    # ==================== Markdown转换实现 ====================
    
    async def _convert_markdown_to_pdf(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str = None) -> str:
        """Markdown转PDF转换"""
        try:
            return await self._markdown_to_pdf_advanced(source_path, file_info, request, trace_id, task_id)
        except Exception as e:
            logger.warning(f"Markdown→PDF高级转换失败: {e}")
            try:
                return await self._markdown_to_pdf_simple(source_path, file_info, request, trace_id, task_id)
            except Exception as e2:
                logger.warning(f"Markdown→PDF简化转换失败: {e2}")
                return await self._markdown_to_pdf_fallback(source_path, file_info, request, trace_id, task_id)
    
    async def _convert_markdown_to_docx(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str = None) -> str:
        """Markdown转DOCX转换"""
        try:
            return await self._markdown_to_docx_advanced(source_path, file_info, request, trace_id, task_id)
        except Exception as e:
            logger.warning(f"Markdown→DOCX高级转换失败: {e}")
            try:
                return await self._markdown_to_docx_simple(source_path, file_info, request, trace_id, task_id)
            except Exception as e2:
                logger.warning(f"Markdown→DOCX简化转换失败: {e2}")
                return await self._markdown_to_docx_fallback(source_path, file_info, request, trace_id, task_id)
    
    async def _markdown_to_pdf_advanced(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str = None) -> str:
        """Markdown转PDF高级实现：完整格式支持"""
        try:
            import markdown2
            import weasyprint
            
            # 读取Markdown文件
            with open(source_path, 'r', encoding='utf-8') as f:
                md_content = f.read()
            
            # 转换为HTML（支持扩展功能）
            html_body = markdown2.markdown(md_content, extras=[
                'fenced-code-blocks', 'tables', 'strike', 'task_list',
                'footnotes', 'header-ids', 'toc'
            ])
            
            # 包装完整HTML
            html_content = f"""
            <!DOCTYPE html>
            <html>
            <head>
                <meta charset="utf-8">
                <title>{request.filename}</title>
                <style>
                    body {{ font-family: 'Segoe UI', Arial, sans-serif; margin: 40px; line-height: 1.6; color: #333; }}
                    h1, h2, h3, h4, h5, h6 {{ color: #2c3e50; margin-top: 20px; }}
                    h1 {{ border-bottom: 2px solid #3498db; padding-bottom: 10px; }}
                    h2 {{ border-bottom: 1px solid #bdc3c7; padding-bottom: 5px; }}
                    code {{ background-color: #f8f9fa; padding: 2px 4px; border-radius: 3px; font-family: 'Courier New', monospace; }}
                    pre {{ background-color: #f8f9fa; padding: 15px; border-radius: 5px; overflow-x: auto; }}
                    blockquote {{ border-left: 4px solid #3498db; margin: 20px 0; padding-left: 20px; color: #555; }}
                    table {{ border-collapse: collapse; width: 100%; margin: 20px 0; }}
                    th, td {{ border: 1px solid #ddd; padding: 12px; text-align: left; }}
                    th {{ background-color: #f2f2f2; font-weight: bold; }}
                    ul, ol {{ margin: 10px 0; padding-left: 30px; }}
                    li {{ margin: 5px 0; }}
                </style>
            </head>
            <body>
                {html_body}
            </body>
            </html>
            """
            
            # 生成输出文件路径
            task_id = task_id or self.file_manager.generate_task_id(DocumentTaskTypePrefix.CONVERT_TASK.value)
            output_filename = f"{request.filename.rsplit('.', 1)[0]}.pdf"
            output_path = self.file_manager.create_file_path(task_id, output_filename)
            
            # 使用weasyprint生成PDF
            # 使用全局进程池和信号量控制并发
            async with _pdf_semaphore:
                loop = asyncio.get_event_loop()
                await loop.run_in_executor(
                    _pdf_executor, 
                    _weasyprint_write_pdf, 
                    html_content, 
                    output_path
                )
            
            # 注册文件到文件管理器
            self.file_manager.register_file(task_id, output_filename, output_path, expire_hours=24)
            
            logger.info(f"Markdown→PDF高级转换成功: {output_path}")
            return output_path
            
        except Exception as e:
            raise Exception(f"高级转换失败: {str(e)}")
    
    async def _markdown_to_pdf_simple(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str = None) -> str:
        """Markdown转PDF简化实现：基础格式"""
        try:
            import markdown2
            import weasyprint
            
            # 读取Markdown文件
            with open(source_path, 'r', encoding='utf-8') as f:
                md_content = f.read()
            
            # 转换为HTML（基础功能）
            html_body = markdown2.markdown(md_content)
            
            # 包装简单HTML
            html_content = f"""
            <!DOCTYPE html>
            <html>
            <head>
                <meta charset="utf-8">
                <title>{request.filename}</title>
                <style>
                    body {{ font-family: Arial, sans-serif; margin: 40px; line-height: 1.6; }}
                    h1, h2, h3 {{ color: #333; }}
                    pre {{ background-color: #f5f5f5; padding: 10px; border-radius: 3px; }}
                    table {{ border-collapse: collapse; width: 100%; margin: 20px 0; }}
                    th, td {{ border: 1px solid #ddd; padding: 8px; text-align: left; }}
                    th {{ background-color: #f2f2f2; }}
                </style>
            </head>
            <body>
                {html_body}
            </body>
            </html>
            """
            
            # 生成输出文件路径
            task_id = task_id or self.file_manager.generate_task_id(DocumentTaskTypePrefix.CONVERT_TASK.value)
            output_filename = f"{request.filename.rsplit('.', 1)[0]}.pdf"
            output_path = self.file_manager.create_file_path(task_id, output_filename)
            
            # 使用weasyprint生成PDF
            # 使用全局进程池和信号量控制并发
            async with _pdf_semaphore:
                loop = asyncio.get_event_loop()
                await loop.run_in_executor(
                    _pdf_executor, 
                    _weasyprint_write_pdf, 
                    html_content, 
                    output_path
                )
            
            # 注册文件到文件管理器
            self.file_manager.register_file(task_id, output_filename, output_path, expire_hours=24)
            
            logger.info(f"Markdown→PDF简化转换成功: {output_path}")
            return output_path
            
        except Exception as e:
            raise Exception(f"简化转换失败: {str(e)}")
    
    async def _markdown_to_pdf_fallback(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str = None) -> str:
        """Markdown转PDF基础实现：纯文本PDF"""
        try:
            import weasyprint
            
            # 读取Markdown文件作为纯文本
            with open(source_path, 'r', encoding='utf-8') as f:
                text_content = f.read()
            
            # 生成简单HTML
            html_content = f"""
            <!DOCTYPE html>
            <html>
            <head>
                <meta charset="utf-8">
                <title>{request.filename}</title>
                <style>
                    body {{ font-family: Arial, sans-serif; margin: 40px; line-height: 1.6; }}
                    h1 {{ color: #333; }}
                    pre {{ white-space: pre-wrap; }}
                </style>
            </head>
            <body>
                <h1>Markdown Content</h1>
                <p><strong>Source:</strong> {request.filename}</p>
                <p><strong>Converted:</strong> {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p>
                <hr>
                <pre>{text_content}</pre>
            </body>
            </html>
            """
            
            # 生成输出文件路径
            task_id = task_id or self.file_manager.generate_task_id(DocumentTaskTypePrefix.CONVERT_TASK.value)
            output_filename = f"{request.filename.rsplit('.', 1)[0]}.pdf"
            output_path = self.file_manager.create_file_path(task_id, output_filename)
            
            # 使用weasyprint生成PDF
            # 使用全局进程池和信号量控制并发
            async with _pdf_semaphore:
                loop = asyncio.get_event_loop()
                await loop.run_in_executor(
                    _pdf_executor, 
                    _weasyprint_write_pdf, 
                    html_content, 
                    output_path
                )
            
            # 注册文件到文件管理器
            self.file_manager.register_file(task_id, output_filename, output_path, expire_hours=24)
            
            logger.info(f"Markdown→PDF基础转换成功: {output_path}")
            return output_path
            
        except Exception as e:
            logger.error(f"基础转换失败: {str(e)}")
            raise Exception("所有转换方法失败")
    
    async def _markdown_to_docx_advanced(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str = None) -> str:
        """Markdown转DOCX高级实现：保持格式"""
        try:
            import markdown2
            from docx import Document
            from docx.shared import Inches
            from bs4 import BeautifulSoup
            
            # 读取Markdown文件
            with open(source_path, 'r', encoding='utf-8') as f:
                md_content = f.read()
            
            # 转换为HTML
            html_content = markdown2.markdown(md_content, extras=['fenced-code-blocks', 'tables'])
            
            # 解析HTML并转换为DOCX
            doc = Document()
            soup = BeautifulSoup(html_content, 'html.parser')
            
            # 处理HTML元素
            for element in soup.find_all(['h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'p', 'pre', 'table']):
                if element.name in ['h1', 'h2', 'h3', 'h4', 'h5', 'h6']:
                    level = int(element.name[1])
                    doc.add_heading(element.get_text(), level=min(level, 3))
                elif element.name == 'p':
                    doc.add_paragraph(element.get_text())
                elif element.name == 'pre':
                    para = doc.add_paragraph(element.get_text())
                    para.style = 'Normal'
                elif element.name == 'table':
                    self._html_table_to_docx(element, doc)
            
            # 生成输出文件路径
            task_id = task_id or self.file_manager.generate_task_id(DocumentTaskTypePrefix.CONVERT_TASK.value)
            output_filename = f"{request.filename.rsplit('.', 1)[0]}.docx"
            output_path = self.file_manager.create_file_path(task_id, output_filename)
            
            # 保存文档
            doc.save(output_path)
            
            # 注册文件到文件管理器
            self.file_manager.register_file(task_id, output_filename, output_path, expire_hours=24)
            
            logger.info(f"Markdown→DOCX高级转换成功: {output_path}")
            return output_path
            
        except Exception as e:
            raise Exception(f"高级转换失败: {str(e)}")
    
    async def _markdown_to_docx_simple(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str = None) -> str:
        """Markdown转DOCX简化实现：基础格式"""
        try:
            import markdown2
            from docx import Document
            from bs4 import BeautifulSoup
            
            # 读取Markdown文件
            with open(source_path, 'r', encoding='utf-8') as f:
                md_content = f.read()
            
            # 转换为HTML（基础功能）
            html_content = markdown2.markdown(md_content)
            
            # 解析HTML并转换为DOCX
            doc = Document()
            soup = BeautifulSoup(html_content, 'html.parser')
            
            # 简单处理HTML元素
            for element in soup.find_all(['h1', 'h2', 'h3', 'p', 'pre']):
                if element.name in ['h1', 'h2', 'h3']:
                    doc.add_heading(element.get_text(), level=1)
                elif element.name in ['p', 'pre']:
                    doc.add_paragraph(element.get_text())
            
            # 生成输出文件路径
            task_id = task_id or self.file_manager.generate_task_id(DocumentTaskTypePrefix.CONVERT_TASK.value)
            output_filename = f"{request.filename.rsplit('.', 1)[0]}.docx"
            output_path = self.file_manager.create_file_path(task_id, output_filename)
            
            # 保存文档
            doc.save(output_path)
            
            # 注册文件到文件管理器
            self.file_manager.register_file(task_id, output_filename, output_path, expire_hours=24)
            
            logger.info(f"Markdown→DOCX简化转换成功: {output_path}")
            return output_path
            
        except Exception as e:
            raise Exception(f"简化转换失败: {str(e)}")
    
    async def _markdown_to_docx_fallback(self, source_path: str, file_info: FileInfo, request: DocumentConvertRequest, trace_id: str, task_id: str = None) -> str:
        """Markdown转DOCX基础实现：纯文本"""
        try:
            from docx import Document
            
            # 读取Markdown文件作为纯文本
            with open(source_path, 'r', encoding='utf-8') as f:
                text_content = f.read()
            
            # 创建Word文档
            doc = Document()
            doc.add_heading(f'Markdown Content: {request.filename}', 0)
            doc.add_paragraph(f'Converted: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')
            doc.add_paragraph(text_content)
            
            # 生成输出文件路径
            task_id = task_id or self.file_manager.generate_task_id(DocumentTaskTypePrefix.CONVERT_TASK.value)
            output_filename = f"{request.filename.rsplit('.', 1)[0]}.docx"
            output_path = self.file_manager.create_file_path(task_id, output_filename)
            
            # 保存文档
            doc.save(output_path)
            
            # 注册文件到文件管理器
            self.file_manager.register_file(task_id, output_filename, output_path, expire_hours=24)
            
            logger.info(f"Markdown→DOCX基础转换成功: {output_path}")
            return output_path
            
        except Exception as e:
            logger.error(f"基础转换失败: {str(e)}")
            raise Exception("所有转换方法失败")
    
    # ==================== DOCX辅助方法 ====================
    
    def _docx_to_html_advanced(self, doc) -> str:
        """DOCX转HTML高级实现：保持格式"""
        html_parts = []
        html_parts.append("""
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="utf-8">
            <style>
                body { font-family: Arial, sans-serif; margin: 40px; line-height: 1.6; }
                h1, h2, h3 { color: #333; }
                table { border-collapse: collapse; width: 100%; margin: 20px 0; }
                th, td { border: 1px solid #ddd; padding: 8px; text-align: left; }
                th { background-color: #f2f2f2; }
                .bold { font-weight: bold; }
                .italic { font-style: italic; }
            </style>
        </head>
        <body>
        """)
        
        # 处理文档元素（段落、图片、表格）
        for element in doc.element.body:
            if element.tag.endswith('p'):  # 段落
                paragraph = None
                for p in doc.paragraphs:
                    if p._element == element:
                        paragraph = p
                        break
                
                if paragraph and paragraph.text.strip():
                    # 检查段落中是否有图片
                    images_html = self._extract_paragraph_images(paragraph)
                    
                    # 检查段落样式
                    if paragraph.style.name.startswith('Heading'):
                        level = min(int(paragraph.style.name[-1]) if paragraph.style.name[-1].isdigit() else 1, 6)
                        html_parts.append(f"<h{level}>{paragraph.text}</h{level}>")
                    else:
                        # 处理段落中的格式
                        para_html = self._process_paragraph_formatting(paragraph)
                        html_parts.append(f"<p>{para_html}</p>")
                    
                    # 添加图片HTML
                    if images_html:
                        html_parts.extend(images_html)
                elif paragraph:
                    # 纯图片段落
                    images_html = self._extract_paragraph_images(paragraph)
                    if images_html:
                        html_parts.extend(images_html)
            
            elif element.tag.endswith('tbl'):  # 表格
                table = None
                for t in doc.tables:
                    if t._element == element:
                        table = t
                        break
                if table:
                    html_parts.append(self._table_to_html(table))
        
        html_parts.append("</body></html>")
        return "\n".join(html_parts)
    
    def _extract_paragraph_images(self, paragraph) -> list:
        """从段落中提取图片并转换为HTML"""
        images_html = []
        try:
            import base64
            from docx.oxml.ns import nsdecls, qn
            
            # 查找段落中的图片
            for run in paragraph.runs:
                for drawing in run._element.xpath('.//w:drawing'):
                    for blip in drawing.xpath('.//a:blip'):
                        # 获取图片关系ID
                        embed_id = blip.get(qn('r:embed'))
                        if embed_id:
                            # 从文档关系中获取图片数据
                            image_part = paragraph.part.related_parts.get(embed_id)
                            if image_part:
                                # 将图片编码为base64
                                image_data = base64.b64encode(image_part.blob).decode()
                                content_type = image_part.content_type
                                
                                # 创建HTML img标签
                                img_html = f'<img src="data:{content_type};base64,{image_data}" style="max-width: 100%; height: auto; margin: 10px 0;" />'
                                images_html.append(img_html)
        
        except Exception as e:
            logger.warning(f"图片提取失败: {str(e)}")
        
        return images_html
    
    def _extract_paragraph_images_to_markdown(self, paragraph) -> list:
        """从段落中提取图片并转换为Markdown格式"""
        images_md = []
        try:
            import base64
            from docx.oxml.ns import nsdecls, qn
            
            # 查找段落中的图片
            for run in paragraph.runs:
                for drawing in run._element.xpath('.//w:drawing'):
                    for blip in drawing.xpath('.//a:blip'):
                        # 获取图片关系ID
                        embed_id = blip.get(qn('r:embed'))
                        if embed_id:
                            # 从文档关系中获取图片数据
                            image_part = paragraph.part.related_parts.get(embed_id)
                            if image_part:
                                # 将图片编码为base64
                                image_data = base64.b64encode(image_part.blob).decode()
                                content_type = image_part.content_type
                                
                                # 创建Markdown图片语法（base64内联）
                                img_md = f'![图片](data:{content_type};base64,{image_data})'
                                images_md.append(img_md)
        
        except Exception as e:
            logger.warning(f"图片提取失败: {str(e)}")
        
        return images_md
    
    def _docx_to_html_simple(self, doc) -> str:
        """DOCX转HTML简化实现：基础格式"""
        html_parts = []
        html_parts.append("""
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="utf-8">
            <style>
                body { font-family: Arial, sans-serif; margin: 40px; line-height: 1.6; }
                h1 { color: #333; }
                p { margin: 10px 0; }
            </style>
        </head>
        <body>
        """)
        
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                html_parts.append(f"<p>{paragraph.text}</p>")
        
        html_parts.append("</body></html>")
        return "\n".join(html_parts)
    
    def _process_paragraph_formatting(self, paragraph) -> str:
        """处理段落中的文本格式"""
        html_text = ""
        for run in paragraph.runs:
            text = run.text
            if run.bold:
                text = f"<strong>{text}</strong>"
            if run.italic:
                text = f"<em>{text}</em>"
            html_text += text
        return html_text
    
    def _table_to_html(self, table) -> str:
        """将DOCX表格转换为HTML"""
        html_parts = ["<table>"]
        
        for i, row in enumerate(table.rows):
            html_parts.append("<tr>")
            for cell in row.cells:
                tag = "th" if i == 0 else "td"
                html_parts.append(f"<{tag}>{cell.text}</{tag}>")
            html_parts.append("</tr>")
        
        html_parts.append("</table>")
        return "\n".join(html_parts)
    
    def _process_paragraph_to_markdown(self, paragraph) -> str:
        """处理段落为Markdown格式"""
        md_text = ""
        for run in paragraph.runs:
            text = run.text
            if run.bold:
                text = f"**{text}**"
            if run.italic:
                text = f"*{text}*"
            md_text += text
        return md_text
    
    def _table_to_markdown(self, table) -> str:
        """将DOCX表格转换为Markdown"""
        markdown_parts = []
        
        for i, row in enumerate(table.rows):
            row_cells = [cell.text.strip() for cell in row.cells]
            markdown_parts.append("| " + " | ".join(row_cells) + " |")
            
            # 添加表头分隔符
            if i == 0:
                separator = "| " + " | ".join(["---"] * len(row_cells)) + " |"
                markdown_parts.append(separator)
        
        return "\n".join(markdown_parts)
    
    # ==================== XLSX辅助方法 ====================
    
    def _xlsx_to_html_advanced(self, wb) -> str:
        """Excel转HTML高级实现"""
        html_parts = []
        html_parts.append("""
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="utf-8">
            <style>
                body { font-family: Arial, sans-serif; margin: 40px; line-height: 1.6; }
                h1, h2 { color: #333; margin-top: 30px; }
                table { border-collapse: collapse; width: 100%; margin: 20px 0; }
                th, td { border: 1px solid #ddd; padding: 8px; text-align: left; }
                th { background-color: #f2f2f2; font-weight: bold; }
                .sheet-title { color: #2c3e50; border-bottom: 2px solid #3498db; padding-bottom: 5px; }
            </style>
        </head>
        <body>
        """)
        
        # 处理每个工作表
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            html_parts.append(f'<h2 class="sheet-title">{sheet_name}</h2>')
            
            # 获取数据范围
            data_rows = list(ws.iter_rows(values_only=True))
            if data_rows:
                html_parts.append('<table>')
                for i, row_data in enumerate(data_rows):
                    html_parts.append('<tr>')
                    tag = 'th' if i == 0 else 'td'
                    for cell_value in row_data:
                        cell_text = str(cell_value) if cell_value is not None else ''
                        html_parts.append(f'<{tag}>{cell_text}</{tag}>')
                    html_parts.append('</tr>')
                html_parts.append('</table>')
            else:
                html_parts.append('<p>No data in this sheet</p>')
        
        html_parts.append('</body></html>')
        return '\n'.join(html_parts)
    
    def _xlsx_to_html_simple(self, wb) -> str:
        """Excel转HTML简化实现"""
        html_parts = []
        html_parts.append("""
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="utf-8">
            <style>
                body { font-family: Arial, sans-serif; margin: 40px; line-height: 1.6; }
                h1 { color: #333; }
                table { border-collapse: collapse; width: 100%; margin: 20px 0; }
                th, td { border: 1px solid #ddd; padding: 8px; text-align: left; }
                th { background-color: #f2f2f2; }
            </style>
        </head>
        <body>
        """)
        
        # 只处理第一个工作表
        ws = wb.active
        data_rows = list(ws.iter_rows(values_only=True))
        
        if data_rows:
            html_parts.append('<table>')
            for i, row_data in enumerate(data_rows):
                html_parts.append('<tr>')
                tag = 'th' if i == 0 else 'td'
                for cell_value in row_data:
                    cell_text = str(cell_value) if cell_value is not None else ''
                    html_parts.append(f'<{tag}>{cell_text}</{tag}>')
                html_parts.append('</tr>')
            html_parts.append('</table>')
        
        html_parts.append('</body></html>')
        return '\n'.join(html_parts)
    
    def _xlsx_to_text(self, wb) -> str:
        """Excel转纯文本"""
        text_parts = []
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            text_parts.append(f"Sheet: {sheet_name}")
            text_parts.append("=" * 50)
            
            data_rows = list(ws.iter_rows(values_only=True))
            for row_data in data_rows:
                row_text = "\t".join([str(cell) if cell is not None else '' for cell in row_data])
                if row_text.strip():
                    text_parts.append(row_text)
            
            text_parts.append("")
        
        return "\n".join(text_parts)
    
    def _excel_data_to_markdown_table(self, data_rows) -> str:
        """Excel数据转Markdown表格"""
        if not data_rows:
            return "No data available"
        
        markdown_parts = []
        
        for i, row_data in enumerate(data_rows):
            row_cells = [str(cell) if cell is not None else '' for cell in row_data]
            markdown_parts.append("| " + " | ".join(row_cells) + " |")
            
            # 添加表头分隔符
            if i == 0:
                separator = "| " + " | ".join(["---"] * len(row_cells)) + " |"
                markdown_parts.append(separator)
        
        return "\n".join(markdown_parts)
    
    def _html_table_to_docx(self, table_element, doc):
        """HTML表格转DOCX表格"""
        rows = table_element.find_all('tr')
        if not rows:
            return
        
        # 创建Word表格
        cols = len(rows[0].find_all(['th', 'td']))
        table = doc.add_table(rows=len(rows), cols=cols)
        table.style = 'Table Grid'
        
        # 填充表格数据
        for i, row in enumerate(rows):
            cells = row.find_all(['th', 'td'])
            for j, cell in enumerate(cells):
                if j < cols:
                    table.cell(i, j).text = cell.get_text().strip()