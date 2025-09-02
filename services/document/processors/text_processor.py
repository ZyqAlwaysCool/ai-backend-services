'''
Description: 文本提取处理器
Author: zyq
Date: 2025-09-02 11:35:42
LastEditors: zyq
LastEditTime: 2025-09-02 11:35:53
'''

import os
from typing import Dict, Any
from datetime import datetime
from loguru import logger

from core.schemas.file_models import FileInfo
from ..schemas import TextExtractRequest, TextExtractResponse
from .base_processor import BaseProcessor


class TextProcessor(BaseProcessor):
    """文本提取处理器"""
    
    def __init__(self, config: Dict[str, Any]):
        super().__init__(config)
    
    async def process(self, request: TextExtractRequest, temp_file_path: str, file_info: FileInfo, trace_id: str = None) -> TextExtractResponse:
        """
        执行DOCX文档文本提取
        
        Args:
            request: 文本提取请求
            temp_file_path: 临时文件路径
            file_info: 文件信息
            trace_id: 链路追踪ID
            
        Returns:
            文本提取响应
        """
        self._log_processing_start("TextProcessor", f"DOCX文本提取", trace_id)
        
        try:
            # 验证文件格式
            if not request.filename.lower().endswith('.docx'):
                raise ValueError("仅支持DOCX格式文件的文本提取")
            
            # 执行文本提取
            text_content = await self._extract_docx_text(temp_file_path, request, trace_id)
            
            # 统计字数
            word_count = len(text_content.replace(' ', '').replace('\n', ''))
            
            # 构建文件元信息
            file_metadata = {
                "filename": request.filename,
                "file_size": file_info.file_size,
                "file_type": file_info.content_type,
                "input_type": request.input_type,
                "extract_options": request.extract_options,
                "processed_at": datetime.now().isoformat()
            }
            
            self._log_processing_end("TextProcessor", "DOCX文本提取", trace_id)
            
            return TextExtractResponse(
                text_content=text_content,
                word_count=word_count,
                file_info=file_metadata
            )
            
        except Exception as e:
            self._log_error("TextProcessor", trace_id, e)
            self._log_processing_end("TextProcessor", "DOCX文本提取", trace_id, success=False)
            raise e
    
    async def _extract_docx_text(self, source_path: str, request: TextExtractRequest, trace_id: str) -> str:
        """从DOCX文件提取文本内容"""
        try:
            from docx import Document
            
            # 读取DOCX文档
            doc = Document(source_path)
            
            # 提取选项
            options = request.extract_options or {}
            include_headers = options.get('include_headers', True)
            include_tables = options.get('include_tables', True)
            include_footnotes = options.get('include_footnotes', False)
            
            text_parts = []
            
            # 提取段落文本
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if text:
                    # 检查是否为标题
                    if paragraph.style.name.startswith('Heading'):
                        if include_headers:
                            text_parts.append(f"\n{text}\n")
                    else:
                        text_parts.append(text)
            
            # 提取表格文本
            if include_tables:
                for table in doc.tables:
                    table_text = self._extract_table_text(table)
                    if table_text:
                        text_parts.append(f"\n[表格内容]\n{table_text}\n")
            
            # 提取脚注文本
            if include_footnotes:
                try:
                    # 尝试提取脚注（如果存在）
                    footnotes = self._extract_footnotes(doc)
                    if footnotes:
                        text_parts.append(f"\n[脚注内容]\n{footnotes}\n")
                except Exception as e:
                    logger.warning(f"脚注提取失败: {e}")
            
            # 合并所有文本
            extracted_text = '\n'.join(text_parts).strip()
            
            if not extracted_text:
                extracted_text = "文档中未找到可提取的文本内容"
            
            logger.info(f"DOCX文本提取成功: 文本长度={len(extracted_text)}")
            return extracted_text
            
        except ImportError:
            raise Exception("缺少python-docx依赖库")
        except Exception as e:
            raise Exception(f"DOCX文本提取失败: {str(e)}")
    
    def _extract_table_text(self, table) -> str:
        """提取表格文本内容"""
        try:
            table_rows = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                if any(row_cells):  # 只添加非空行
                    table_rows.append('\t'.join(row_cells))
            
            return '\n'.join(table_rows)
        except Exception as e:
            logger.warning(f"表格文本提取失败: {e}")
            return ""
    
    def _extract_footnotes(self, doc) -> str:
        """提取脚注内容"""
        try:
            footnotes = []
            # 这里可以添加脚注提取逻辑
            # python-docx对脚注的支持有限，可以考虑使用其他方法
            return '\n'.join(footnotes)
        except Exception as e:
            logger.warning(f"脚注提取失败: {e}")
            return ""